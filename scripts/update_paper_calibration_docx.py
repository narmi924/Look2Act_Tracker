"""Apply targeted thesis DOCX updates for calibration and resolution flow."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "final-project-paper" / "论文" / "实时视线追踪技术研究与实现.docx"
IMG = ROOT / "docs" / "final-project-paper" / "引用图片"


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text.strip()


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph_text(paragraph) == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_paragraph_contains(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if text in paragraph_text(paragraph):
            return paragraph
    raise ValueError(f"Paragraph containing text not found: {text}")


def set_paragraph_text(doc: Document, old_text: str, new_text: str) -> None:
    paragraph = find_paragraph(doc, old_text)
    paragraph.clear()
    paragraph.add_run(new_text)


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def image_paragraph_before_caption(doc: Document, caption: str) -> Paragraph:
    caption_p = find_paragraph(doc, caption)
    caption_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is caption_p._p
    )
    for i in range(caption_index - 1, max(-1, caption_index - 8), -1):
        paragraph = doc.paragraphs[i]
        if paragraph._p.xpath(".//w:drawing") or paragraph._p.xpath(".//w:pict"):
            return paragraph
    raise ValueError(f"Image paragraph not found before caption: {caption}")


def replace_image_before_caption(doc: Document, caption: str, image_path: Path, width_cm: float) -> None:
    paragraph = image_paragraph_before_caption(doc, caption)
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))


def format_table(table: Table, font_size: float = 8.0) -> None:
    table.style = "Table Grid"
    table.autofit = True
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    if row_index == 0:
                        run.bold = True


def fill_table(table: Table, rows: list[list[str]], font_size: float = 8.0) -> None:
    target_cols = len(rows[0])
    while len(table.columns) < target_cols:
        table.add_column(Cm(2.0))
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)

    for row_index, row_data in enumerate(rows):
        row = table.rows[row_index]
        for col_index, value in enumerate(row_data):
            row.cells[col_index].text = value
    format_table(table, font_size=font_size)


def apply_chapter_5_updates(doc: Document) -> None:
    set_paragraph_text(
        doc,
        "Look2Act 采用桌面图形界面组织系统功能，主要页面包括主页、摄像头预览页、校准页、实时追踪页、设置页、全屏验证窗口和全屏交互窗口。用户按照“启动—预览—校准—追踪—验证—交互”的流程使用系统。",
        "Look2Act 采用桌面图形界面组织系统功能，主要页面包括主页、摄像头预览页、校准页、实时追踪页、设置页、全屏验证窗口和全屏交互窗口。用户按照“启动—模式选择—摄像头预览—分辨率与显示环境确认—25 点校准—实时追踪—全屏验证—注视交互”的流程使用系统。",
    )
    set_paragraph_text(
        doc,
        "摄像头预览页用于检查摄像头是否可用、人脸是否能够被稳定检测、眼区裁剪是否正常。用户在进入校准前可以通过预览页调整坐姿、摄像头角度、光照和屏幕位置。该页面是减少后续校准失败的重要步骤。",
        "摄像头预览页用于检查摄像头是否可用、人脸是否能够被稳定检测、眼区裁剪是否正常。用户在进入校准前可以通过预览页调整坐姿、摄像头角度、光照和屏幕位置，并确认摄像头请求分辨率与实际画面状态。该页面是减少后续校准失败的重要步骤。",
    )

    camera_heading = find_paragraph(doc, "5.9.2 摄像头预览")
    camera_para = find_paragraph_contains(doc, "摄像头预览页用于检查摄像头是否可用")
    inserted_heading = insert_paragraph_after(camera_para, "5.9.3 分辨率与显示环境确认", style=camera_heading.style.name)
    insert_paragraph_after(
        inserted_heading,
        "进入校准前，系统读取当前主屏幕几何尺寸，并据此生成校准目标点和后续像素误差计算所需的屏幕坐标。用户需要在预览或设置阶段确认摄像头输入、请求分辨率、人脸检测、眼区裁剪和显示环境是否稳定。屏幕分辨率影响校准点位置、目标点坐标和像素误差；摄像头分辨率影响人脸检测质量、眼区裁剪质量、帧率和实时稳定性。系统中的 128×128 表示模型眼区输入尺寸，1280×720 表示摄像头请求或预览分辨率，二者均不是固定屏幕分辨率。",
    )

    set_paragraph_text(doc, "5.9.3 校准页面", "5.9.4 校准页面")
    set_paragraph_text(doc, "5.9.4 实时追踪页面", "5.9.5 实时追踪页面")
    set_paragraph_text(doc, "5.9.5 设置页面", "5.9.6 设置页面")

    set_paragraph_text(
        doc,
        "系统使用流程如图5-12所示。",
        "系统使用流程如图5-12所示，其中分辨率与显示环境确认位于摄像头预览之后、25 点校准之前。",
    )
    replace_image_before_caption(doc, "图5-12 Look2Act 使用流程", IMG / "5-12.png", 14.5)


def apply_chapter_6_updates(doc: Document) -> None:
    set_paragraph_text(doc, "6.5 校准映射实验", "6.5 校准映射实验与最终25点校准验证")
    set_paragraph_text(
        doc,
        "本文对不同校准点数量和映射方式进行了对比分析。校准实验的目的不是替代最终系统中的 25 点实时校准流程，而是分析少量校准点条件下不同映射策略的稳定性，为校准模块设计提供依据。实验比较了无校准、少点校准、仿射映射和多项式映射等方式对屏幕误差的影响。",
        "本文对不同校准点数量和映射方式进行了对比分析。本节区分两类结果：第一类是离线校准对比，评价方式为 hold-out 像素误差，用于比较不同校准点数和映射方式的泛化表现；第二类是最终实时校准结果，评价方式为校准拟合残差，用于记录 Classic 与 Deep 在最终 25 点 polynomial 配置下的运行期校准文件。两类指标来源和含义不同，不能直接等同。",
    )
    set_paragraph_text(
        doc,
        "少点校准对比结果见表6-4。",
        "校准策略对比与最终 25 点校准结果见表6-4。表中“离线校准对比”为 hold-out 像素误差，“最终实时校准”为校准拟合残差。",
    )
    set_paragraph_text(doc, "表6-4 少点校准映射对比结果", "表6-4 校准策略对比与最终25点校准结果")
    set_paragraph_text(doc, "图6-6 不同校准方式下的平均像素误差", "图6-6 不同校准策略与最终25点校准结果对比")
    replace_image_before_caption(doc, "图6-6 不同校准策略与最终25点校准结果对比", IMG / "6-6.png", 15.0)

    set_paragraph_text(
        doc,
        "从表6-4可以看出，9点仿射校准能够将平均像素误差从 314.5 px 降低到 221.2 px，说明少量校准点已经能够有效修正部分系统性偏差。相比之下，9点多项式映射误差反而升高到 437.0 px，说明在校准点数量较少时，高阶映射容易受到噪声样本影响，出现过拟合和泛化不稳定问题。",
        "从少点校准结果可以看出，9 点 affine 校准能够将 hold-out 平均像素误差从 314.54 px 降低到 221.19 px，说明少量校准点已经能够修正部分系统性偏差。相比之下，9 点 polynomial 的 hold-out 平均误差升高到 436.97 px，说明在校准点较少时，高阶映射容易受到噪声样本影响，出现过拟合和泛化不稳定问题。",
    )
    set_paragraph_text(
        doc,
        "该结果说明，校准映射不是越复杂越好。对于少点校准，稳定性和泛化能力比表达能力更重要。仿射映射虽然形式简单，但参数较少，对噪声更不敏感，因此在少量校准点条件下表现更稳定。",
        "扩展校准实验进一步表明，增加到 25 点后校准误差和稳定性整体优于 9 点策略。25 点 affine 在本轮 hold-out 矩阵中平均误差最低，为 184.57 px；25 点 polynomial 的平均误差为 198.49 px，虽不是绝对最低，但明显优于 9 点 affine 和 9 点 polynomial。该结果说明，polynomial 映射需要更充分的屏幕覆盖，25 点覆盖屏幕中心、边缘和角落后，其稳定性明显改善。",
    )
    set_paragraph_text(
        doc,
        "最终 Look2Act 系统采用 5×5 共25点的实时校准流程。25点校准能够覆盖屏幕中心、边缘和角落区域，为映射函数提供更完整的空间约束。在最终软件系统中，Classic 模式和 Deep 模式均可通过25点校准采集原始点与目标点对应关系，并拟合校准映射用于后续追踪、验证和交互。",
        "最终 Look2Act 系统采用 5×5 共 25 点的实时校准流程。考虑到最终系统 Classic 与 Deep 均配置为 25 点 polynomial，且 25 点覆盖屏幕中心、边缘和角落，最终系统采用 25 点 polynomial 作为实时校准配置。该选择由扩展实验结果、屏幕覆盖完整性和系统配置一致性共同决定，而不是因为 25 点 polynomial 在所有实验指标中绝对最优。",
    )
    set_paragraph_text(
        doc,
        "需要区分的是，表6-4中的9点仿射校准结果属于少点校准对比实验，用于分析校准点数量和映射函数选择对误差的影响；最终软件演示中的25点校准是实时系统配置，用于提高真实交互阶段的屏幕覆盖范围和校准稳定性。二者服务于不同目的，不能简单等同。",
        "需要区分的是，离线校准对比中的数值为 hold-out 像素误差，用于衡量校准映射在未参与拟合样本上的泛化表现；运行期 Classic 和 Deep 的 25 点结果来自真实校准文件，只能表述为校准拟合残差，不能写成验证误差或泛化误差。当前运行期 Classic 25 点 polynomial 的平均拟合残差为 73.58 px，最大拟合残差为 212.54 px；Deep 25 点 polynomial 的平均拟合残差为 118.69 px，最大拟合残差为 251.97 px。最终系统配置汇总见表6-5。",
    )
    set_paragraph_text(doc, "表6-5 最终系统25点校准验证结果", "表6-5 最终系统25点实时校准拟合残差")

    table_64_rows = [
        ["实验类型", "点数", "映射方式", "评价方式", "平均误差/残差", "标准差或最大值", "说明"],
        ["离线校准对比", "0", "无校准", "hold-out 像素误差", "314.54 px", "-", "原始屏幕映射基线"],
        ["离线校准对比", "9", "affine", "hold-out 像素误差", "221.19 px", "28.59 px", "少点校准中较稳定"],
        ["离线校准对比", "9", "polynomial", "hold-out 像素误差", "436.97 px", "222.27 px", "少点条件下不稳定"],
        ["离线校准对比", "25", "affine", "hold-out 像素误差", "184.57 px", "12.52 px", "本轮 hold-out 平均误差最低"],
        ["离线校准对比", "25", "polynomial", "hold-out 像素误差", "198.49 px", "16.73 px", "优于 9 点策略，作为最终配置参考"],
        ["最终实时校准", "25", "polynomial", "校准拟合残差", "Classic 73.58 px", "最大 212.54 px", "运行期 calibration_classic.json"],
        ["最终实时校准", "25", "polynomial", "校准拟合残差", "Deep 118.69 px", "最大 251.97 px", "运行期 calibration_deep.json"],
    ]
    fill_table(doc.tables[89], table_64_rows, font_size=7.5)

    table_65_rows = [
        ["模式", "校准点数", "映射方式", "评价方式", "平均拟合残差", "中位拟合残差", "最大拟合残差"],
        ["Classic", "25", "polynomial", "校准拟合残差", "73.58 px", "53.96 px", "212.54 px"],
        ["Deep", "25", "polynomial", "校准拟合残差", "118.69 px", "103.46 px", "251.97 px"],
    ]
    fill_table(doc.tables[90], table_65_rows, font_size=8.0)

    set_paragraph_text(
        doc,
        "在校准实验中，9点仿射校准能够将平均像素误差从 314.5 px 降低到 221.2 px，而9点多项式映射出现误差升高，说明少点校准场景中稳定性比映射复杂度更重要。最终系统采用25点实时校准流程，用于覆盖更完整的屏幕区域并服务于真实交互阶段。头姿消融实验表明，头姿相关信息对屏幕误差具有重要影响，但在普通摄像头实时系统中，头姿估计、模型输出和屏幕几何之间的坐标约定必须严格统一。",
        "在校准实验中，9 点 affine 结果说明少点校准时参数较少的映射方式更稳定；扩展到 25 点后，校准误差和稳定性整体优于 9 点策略，其中 25 点 affine 在本轮 hold-out 矩阵中平均误差最低，25 点 polynomial 则在充分屏幕覆盖后明显优于 9 点策略。最终系统采用 25 点 polynomial，是实验结果、屏幕覆盖和 Classic/Deep 系统配置一致性共同决定的。头姿消融实验表明，头姿相关信息对屏幕误差具有重要影响，但在普通摄像头实时系统中，头姿估计、模型输出和屏幕几何之间的坐标约定必须严格统一。",
    )
    set_paragraph_text(
        doc,
        "第四，用户校准能有效修正系统性偏差，少点校准对比实验表明，9点affine校准已能把平均像素误差从314.5 px降低到221.2 px，说明少量校准点就足以显著修正系统偏移；而9点polynomial映射的误差则升高到437.0 px，说明校准点较少时高阶映射容易受噪声干扰而出现过拟合现象，我最终把系统设计为采用25点实时校准流程，借助更完整的屏幕覆盖来提升真实交互阶段的校准稳定性。",
        "第四，用户校准能有效修正系统性偏差。少点校准对比实验表明，9 点 affine 校准可将 hold-out 平均像素误差从 314.54 px 降低到 221.19 px，而 9 点 polynomial 映射误差升高到 436.97 px，说明少点条件下 affine 更稳定。扩展实验进一步表明，25 点策略整体优于 9 点策略，其中 25 点 affine 的 hold-out 平均误差最低，25 点 polynomial 虽不是绝对最低，但在充分屏幕覆盖后明显优于 9 点 affine 和 9 点 polynomial。最终选择 25 点 polynomial 是实验结果、屏幕覆盖和系统配置一致性共同决定的。",
    )


def apply_appendix_b_updates(doc: Document) -> None:
    preview_caption = find_paragraph(doc, "图 B-2 摄像头预览界面")
    b4_heading = insert_paragraph_after(preview_caption, "B.4 确认分辨率", style="Heading 3")
    b4_body = insert_paragraph_after(
        b4_heading,
        "在进入 25 点校准前，用户需要确认屏幕分辨率、摄像头输入和显示环境。系统根据当前主屏幕几何尺寸生成校准目标点；设置页面可用于确认摄像头请求分辨率、屏幕物理尺寸、屏幕距离和校准点数等参数。屏幕分辨率影响目标点像素坐标和误差计算，摄像头分辨率影响检测质量、眼区裁剪和实时帧率。",
    )
    image_p = insert_paragraph_after(b4_body)
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.add_run().add_picture(str(IMG / "B-3.png"), width=Cm(13.5))
    insert_paragraph_after(image_p, "图 B-3 分辨率与系统设置界面")

    set_paragraph_text(doc, "B.4 25 点校准", "B.5 25 点校准")
    set_paragraph_text(doc, "图 B-3 25 点校准界面", "图 B-4 25 点校准界面")
    set_paragraph_text(doc, "B.5 实时追踪与全屏验证", "B.6 实时追踪与全屏验证")
    set_paragraph_text(doc, "图 B-4 实时追踪界面", "图 B-5 实时追踪界面")
    set_paragraph_text(doc, "图 B-5 全屏验证界面", "图 B-6 全屏验证界面")
    set_paragraph_text(doc, "B.6 注视交互演示", "B.7 注视交互演示")
    set_paragraph_text(doc, "图 B-6 交互启动器界面", "图 B-7 交互启动器界面")
    set_paragraph_text(doc, "图 B-7 五子棋注视交互界面", "图 B-8 五子棋注视交互界面")
    set_paragraph_text(doc, "B.7 常见问题与处理方式", "B.8 常见问题与处理方式")

    replace_image_before_caption(doc, "图 B-4 25 点校准界面", IMG / "B-4.png", 13.5)
    replace_image_before_caption(doc, "图 B-5 实时追踪界面", IMG / "B-5.png", 13.5)
    replace_image_before_caption(doc, "图 B-6 全屏验证界面", IMG / "B-6.png", 13.5)
    replace_image_before_caption(doc, "图 B-7 交互启动器界面", IMG / "B-7.png", 13.5)
    replace_image_before_caption(doc, "图 B-8 五子棋注视交互界面", IMG / "B-8.png", 13.5)


def main() -> None:
    doc = Document(DOCX)
    apply_chapter_5_updates(doc)
    apply_chapter_6_updates(doc)
    apply_appendix_b_updates(doc)
    doc.save(DOCX)


if __name__ == "__main__":
    main()
